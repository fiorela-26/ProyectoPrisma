import pymupdf
import pytest

from services import ai_router, extractor


@pytest.fixture(autouse=True)
def sin_cache(monkeypatch):
    monkeypatch.setattr("services.extractor.cache_ia.buscar", lambda *a, **k: None)
    monkeypatch.setattr("services.extractor.cache_ia.guardar", lambda *a, **k: None)


def _pdf_con_texto(paginas: int = 2) -> bytes:
    doc = pymupdf.open()
    for i in range(paginas):
        pagina = doc.new_page()
        pagina.insert_text((50, 100), f"Página {i + 1}. " + "Texto educativo de prueba. " * 10, fontsize=11)
    return doc.tobytes()


def _pdf_sin_texto(paginas: int = 1) -> bytes:
    doc = pymupdf.open()
    for _ in range(paginas):
        pagina = doc.new_page()
        pagina.draw_rect(pymupdf.Rect(20, 20, 200, 200), color=(0, 0, 0), fill=(0.5, 0.5, 0.5))
    return doc.tobytes()


def test_pdf_con_texto_usa_pymupdf_sin_ia(monkeypatch):
    def no_debe_llamarse(*a, **k):
        raise AssertionError("No debe llamarse a la IA para un PDF con texto")

    monkeypatch.setattr(ai_router, "gemini_generar", no_debe_llamarse)
    res = extractor.extraer(_pdf_con_texto(2), "pdf")
    assert res.metodo == "pymupdf"
    assert res.paginas == 2
    assert "Página 1" in res.texto and "Página 2" in res.texto
    assert "<!-- página 2 -->" in res.texto


def test_pdf_escaneado_sin_tesseract_usa_gemini(monkeypatch):
    llamadas = []

    def gemini_falso(partes, cadena=None, temperatura=0.1):
        llamadas.append(partes)
        return ai_router.RespuestaIA(texto="# Título OCR\n\nTexto reconocido.", modelo="modelo-vision-x")

    monkeypatch.setattr(extractor, "_tesseract_disponible", lambda: False)
    monkeypatch.setattr(ai_router, "gemini_generar", gemini_falso)
    res = extractor.extraer(_pdf_sin_texto(), "pdf")
    assert res.metodo == "gemini"
    assert res.modelo_usado == "modelo-vision-x"
    assert res.texto.startswith("# Título OCR")
    assert len(llamadas) == 1
    assert extractor.PROMPT_OCR in llamadas[0]
    assert "DATOS" in extractor.PROMPT_OCR  # el documento se marca como datos, no instrucciones


def test_texto_plano():
    res = extractor.extraer("Hola mundo\n".encode("utf-8"), "texto")
    assert res.metodo == "texto"
    assert res.texto == "Hola mundo"


def test_texto_latin1_no_rompe():
    res = extractor.extraer("Año de educación".encode("latin-1"), "texto")
    assert "educaci" in res.texto


def _docx_bytes(texto: str) -> bytes:
    import io

    from docx import Document

    doc = Document()
    doc.add_heading("Título de clase", level=1)
    doc.add_paragraph(texto)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_bytes(texto: str) -> bytes:
    import io

    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    caja = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    caja.text_frame.text = texto
    caja.text_frame.paragraphs[0].font.size = Pt(18)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _xlsx_bytes(texto: str) -> bytes:
    import io

    from openpyxl import Workbook

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Hoja1"
    hoja["A1"] = texto
    hoja["B1"] = 12
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_docx_extrae_texto():
    res = extractor.extraer(_docx_bytes("La célula es la unidad de la vida."), "word")
    assert res.metodo == "docx"
    assert "célula" in res.texto
    assert "Título de clase" in res.texto


def test_pptx_extrae_texto():
    res = extractor.extraer(_pptx_bytes("Fotosíntesis en plantas"), "presentacion")
    assert res.metodo == "pptx"
    assert "Fotosíntesis" in res.texto
    assert "Diapositiva 1" in res.texto


def test_xlsx_extrae_celdas():
    res = extractor.extraer(_xlsx_bytes("Energía"), "hoja")
    assert res.metodo == "xlsx"
    assert "Energía" in res.texto
    assert "12" in res.texto


def test_proponer_titulo_usa_encabezado_sin_ia(monkeypatch):
    def no_ia(*a, **k):
        raise AssertionError("No debe llamar a la IA si ya hay un encabezado")

    monkeypatch.setattr(ai_router, "groq_texto", no_ia)
    titulo = extractor.proponer_titulo("# La fotosíntesis\n\nLas plantas convierten luz en energía.", "archivo.pdf")
    assert titulo == "La fotosíntesis"


def test_proponer_titulo_usa_ia_si_no_hay_encabezado(monkeypatch):
    monkeypatch.setattr(
        ai_router,
        "groq_texto",
        lambda *a, **k: ai_router.RespuestaIA(texto='{"titulo": "La célula"}', modelo="lite"),
    )
    titulo = extractor.proponer_titulo("La célula es la unidad estructural de los seres vivos.", "clase.pdf")
    assert titulo == "La célula"


def test_filtrar_relleno_quita_indice():
    crudo = "Índice\n1. Derivada ........ 3\nPágina 1\n<!-- página 2 -->\nLa derivada mide el cambio instantáneo de una función."
    limpio = extractor._filtrar_relleno(crudo)
    assert "Derivada ........ 3" not in limpio
    assert "cambio instantáneo" in limpio or "cambio instantaneo" in limpio.replace("á", "a")


def test_preparar_clase_rechaza_ia_con_indice(monkeypatch):
    monkeypatch.setattr(
        ai_router,
        "groq_texto",
        lambda *a, **k: ai_router.RespuestaIA(
            texto='{"titulo": "Curso", "texto": "## Resumen\\nÍndice\\n1. Tema ........ 2\\n## Ideas principales\\n- Nada"}',
            modelo="lite",
        ),
    )
    crudo = "Índice\n1. Tema ........ 2\n\nLa derivada es el límite del cociente incremental.\n\nSirve para calcular la pendiente de una curva."
    out = extractor.preparar_clase(crudo, fallback_titulo="Derivadas")
    assert "## Resumen" in out["texto"]
    assert "........" not in out["texto"]
    assert "Índice" not in out["texto"]


def test_preparar_clase_si_falla_ia_no_devuelve_indice(monkeypatch):
    def boom(*a, **k):
        raise RuntimeError("sin ia")

    monkeypatch.setattr(ai_router, "groq_texto", boom)
    crudo = "Índice\nTema 1 ........ 2\nLa derivada es el límite del cociente incremental.\n\nSirve para calcular la pendiente de una curva."
    out = extractor.preparar_clase(crudo, fallback_titulo="Derivadas")
    assert "## Resumen" in out["texto"]
    assert "........" not in out["texto"]
    assert "Índice" not in out["texto"]


def test_preparar_clase_arma_resumen_y_ejemplos(monkeypatch):
    monkeypatch.setattr(
        ai_router,
        "groq_texto",
        lambda *a, **k: ai_router.RespuestaIA(
            texto='{"titulo": "Derivadas", "texto": "## Resumen\\nCambio instantaneo.\\n## Ideas principales\\n- Limite del cociente.\\n## Explicacion del tema\\nLa derivada mide la pendiente.\\n## Ejemplos\\nSi f(x)=x^2, la derivada es 2x."}',
            modelo="lite",
        ),
    )
    out = extractor.preparar_clase("f'(x) = lim h->0 [f(x+h)-f(x)]/h", fallback_titulo="archivo")
    assert out["titulo"] == "Derivadas"
    assert "## Resumen" in out["texto"]
    assert "## Ejemplos" in out["texto"]


def test_preparar_clase_incluye_nivel_cine_en_el_prompt(monkeypatch):
    visto = {}

    def fake_groq(mensajes, **k):
        visto["mensajes"] = mensajes
        return ai_router.RespuestaIA(
            texto='{"titulo": "Derivadas", "texto": "## Resumen\\nCambio instantaneo.\\n## Ideas principales\\n- Limite.\\n## Explicacion del tema\\nPendiente.\\n## Ejemplos\\n2x."}',
            modelo="lite",
        )

    monkeypatch.setattr(ai_router, "groq_texto", fake_groq)
    nivel = "CINE 2 - Secundaria inferior / Ciclo básico (12 a 14 años)"
    extractor.preparar_clase("La derivada mide el cambio instantáneo.", fallback_titulo="archivo", nivel_educativo=nivel)
    system = visto["mensajes"][0]["content"]
    user = visto["mensajes"][1]["content"]
    assert "UNESCO (CINE)" in system
    assert nivel in system
    assert "Lenguaje accesible" in system
    assert "Tono dinámico" in system
    assert "Enfoque crítico" in system
    assert "Rigor académico" in system
    assert nivel in user


def test_generar_clase_desde_titulo_incluye_tema_y_cine(monkeypatch):
    visto = {}

    def fake_groq(mensajes, **k):
        visto["mensajes"] = mensajes
        return ai_router.RespuestaIA(
            texto='{"titulo": "La célula", "texto": "## Resumen\\nUnidad de la vida.\\n## Ideas principales\\n- Membrana.\\n## Explicacion del tema\\nLa célula es la unidad.\\n## Ejemplos\\nUn glóbulo rojo."}',
            modelo="lite",
        )

    monkeypatch.setattr(ai_router, "groq_texto", fake_groq)
    nivel = "CINE 1 - Primaria inicial (6 a 8 años)"
    out = extractor.generar_clase_desde_titulo("La célula", nivel_educativo=nivel)
    system = visto["mensajes"][0]["content"]
    user = visto["mensajes"][1]["content"]
    assert out["titulo"] == "La célula"
    assert "## Resumen" in out["texto"]
    assert "UNESCO (CINE)" in system
    assert nivel in system
    assert "La célula" in user
    assert "No hay documento adjunto" in user


def test_depurar_contenido_quita_relleno(monkeypatch):
    monkeypatch.setattr(
        ai_router,
        "groq_texto",
        lambda *a, **k: ai_router.RespuestaIA(
            texto='{"texto": "## La célula\\n\\nUnidad estructural de los seres vivos."}',
            modelo="lite",
        ),
    )
    limpio = extractor.depurar_contenido("Diapositiva 1\nClic para agregar título\nCopyright 2020\nLa célula es la unidad estructural de los seres vivos.")
    assert "célula" in limpio.lower()
    assert "Clic para agregar" not in limpio


def test_audio_no_soportado_aqui():
    with pytest.raises(extractor.ExtraccionNoSoportada):
        extractor.extraer(b"xxx", "audio")


def test_router_pasa_al_siguiente_modelo_en_429(monkeypatch):
    class Err429(ai_router.genai_errors.APIError):
        def __init__(self):
            self.code = 429
            self.message = "quota"

    intentos = []

    class ClienteFalso:
        class models:  # noqa: N801
            @staticmethod
            def generate_content(model, contents, config):
                intentos.append(model)
                if model == "primario":
                    raise Err429()

                class R:
                    text = "ok"

                return R()

    monkeypatch.setattr(ai_router, "_gemini", lambda: ClienteFalso())
    monkeypatch.setattr(ai_router.time, "sleep", lambda s: None)
    res = ai_router.gemini_generar(["hola"], cadena=("primario", "respaldo"))
    assert res.modelo == "respaldo"
    assert intentos == ["primario", "respaldo"]
